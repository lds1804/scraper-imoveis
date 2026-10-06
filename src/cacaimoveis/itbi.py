"""Baixa a base de ITBI da Prefeitura de São Paulo.

Fonte (descoberta em 2026-10, com ajuda do usuário):
    https://prefeitura.sp.gov.br/web/fazenda/w/acesso_a_informacao/31501

O que é: "Dados das Transações Imobiliárias com recolhimento de ITBI". Cada
linha é uma Declaração de Transações Imobiliárias (DTI) efetivamente paga no
mês de referência. É o registro OFICIAL do preço declarado de venda de cada
imóvel de São Paulo — o dado que nenhum portal de anúncio tem.

Por que é valioso para este projeto: dá para comparar o preço PEDIDO no
anúncio com o preço REAL de transações na mesma rua/bairro, além do valor
venal de referência da Prefeitura.

Armadilha importante: os arquivos NÃO ficam em dados abertos (o portal CKAN
da Prefeitura não tem nada de ITBI — verificado, 0 resultados em 488
datasets). Ficam nesta página de "Acesso à Informação", com URL que muda a
cada atualização mensal. Por isso o script RASPA a página em vez de fixar
links.

Uso:
    .venv\\Scripts\\python.exe itbi.py --baixar            # todos os anos
    .venv\\Scripts\\python.exe itbi.py --baixar --anos 2024 2025 2026
    .venv\\Scripts\\python.exe itbi.py --listar             # só mostra os links
    .venv\\Scripts\\python.exe itbi.py --inspecionar        # estrutura do arquivo
"""

from __future__ import annotations

import argparse
import os
import re
import sys
import time

import requests

from cacaimoveis import config, logs

log = logs.obter(__name__)

PAGINA = "https://prefeitura.sp.gov.br/web/fazenda/w/acesso_a_informacao/31501"
ITBI_DIR = config.caminho("dados", "itbi")

HEADERS = {
    "User-Agent": config.USER_AGENT,
    "Accept": "text/html,application/xhtml+xml,*/*",
}


# ---------------------------------------------------------------------------
# Descoberta dos links
# ---------------------------------------------------------------------------
def descobrir_links(verbose: bool = True) -> dict[int, str]:
    """Raspa a página e devolve {ano: url_do_xlsx}.

    Os links mudam a cada atualização mensal (o nome do arquivo carrega a data
    de extração, ex. `GUIAS DE ITBI PAGAS (30092026).xlsx`), então fixar as
    URLs no código quebraria todo mês.

    O layout da página é uma lista onde cada item tem o ano seguido dos links:
        2024 (<a href="...xlsx">Excel/xlsx</a>) (<a href="...ods">ODS</a>)

    CUIDADO com duas armadilhas reais:
      1. Os links são HTML (`<a href="...">`), não markdown. E o mesmo arquivo
         aparece com e sem `www.` — daí usar `re.findall` em todo o bloco.
      2. Os nomes de arquivo CONTÊM parênteses
         (`GUIAS_DE_ITBI_PAGAS_(2019).xlsx`), então não dá para delimitar o
         bloco por `[^)]*` — ele cortaria no primeiro `)`. O bloco é delimitado
         pelo ANO SEGUINTE, não por parêntese.
    """
    r = requests.get(PAGINA, headers=HEADERS, timeout=90)
    r.raise_for_status()
    html = r.text

    # âncoras de ano: "2024 (" — cada um marca o começo de um item da lista.
    marcas = [(m.start(), int(m.group(1)))
              for m in re.finditer(r"\b(20\d{2})\s*\(", html)]

    anos: dict[int, str] = {}
    for i, (pos, ano) in enumerate(marcas):
        # o bloco vai do ano até o próximo ano (ou o fim, para o último)
        fim = marcas[i + 1][0] if i + 1 < len(marcas) else len(html)
        bloco = html[pos:fim]

        urls = re.findall(r'href="([^"]+\.xlsx)"', bloco, re.I)
        if not urls:
            continue
        # mantém a primeira ocorrência do ano (a lista é cronológica inversa,
        # então um ano pode reaparecer em texto explicativo depois)
        if ano not in anos:
            anos[ano] = urls[0]

    if verbose:
        print(f"Página: {PAGINA}")
        if anos:
            print(f"Anos encontrados: {len(anos)} ({min(anos)}–{max(anos)})\n")
        else:
            print("Nenhum ano encontrado.\n")
    return anos


# ---------------------------------------------------------------------------
# Download
# ---------------------------------------------------------------------------
def _nome_local(ano: int, url: str) -> str:
    return os.path.join(ITBI_DIR, f"itbi_{ano}.xlsx")


def _baixar_um(ano: int, url: str, forcar: bool = False) -> tuple[str, int]:
    """Baixa um ano. Devolve (situação, bytes)."""
    destino = _nome_local(ano, url)
    if os.path.exists(destino) and not forcar:
        tam = os.path.getsize(destino)
        return "já tinha", tam

    with requests.get(url, headers=HEADERS, timeout=300, stream=True) as r:
        r.raise_for_status()
        tmp = destino + ".parcial"
        with open(tmp, "wb") as f:
            for bloco in r.iter_content(1 << 16):
                f.write(bloco)
        os.replace(tmp, destino)   # troca atômica: nunca deixa arquivo pela metade

    return "baixado", os.path.getsize(destino)


def baixar(anos_pedidos: list[int] | None, forcar: bool = False) -> int:
    os.makedirs(ITBI_DIR, exist_ok=True)
    links = descobrir_links()
    if not links:
        print("Nenhum link encontrado. O layout da página mudou?")
        return 1

    alvos = []
    if anos_pedidos:
        for a in anos_pedidos:
            if a in links:
                alvos.append((a, links[a]))
            else:
                print(f"[aviso] {a} não está na página (anos: "
                      f"{min(links)}–{max(links)})")
    else:
        alvos = sorted(links.items())

    if not alvos:
        return 1

    print(f"Baixando {len(alvos)} arquivo(s) para {os.path.relpath(ITBI_DIR, config.RAIZ)}/")
    print(f"  {'ano':>5s}  {'situação':>10s}  {'tamanho':>10s}  tempo")
    print("  " + "-" * 46)

    total_bytes = 0
    t_inicio = time.time()
    for ano, url in alvos:
        t0 = time.time()
        try:
            situacao, tam = _baixar_um(ano, url, forcar)
        except Exception as e:  # noqa: BLE001
            log.warning("erro tratado, a execução segue: %s", e, exc_info=True)
            print(f"  {ano:>5d}  {'ERRO':>10s}  {str(e)[:34]}")
            continue
        total_bytes += tam
        print(f"  {ano:>5d}  {situacao:>10s}  {tam/1e6:>8.1f} MB  "
              f"{time.time()-t0:>5.1f}s")

    print("  " + "-" * 46)
    print(f"  total: {total_bytes/1e9:.2f} GB em {time.time()-t_inicio:.0f}s")
    return 0


# ---------------------------------------------------------------------------
# Inspeção (para conferir o layout antes de escrever o parser)
# ---------------------------------------------------------------------------
def inspecionar(ano: int, aba: str | None = None, linhas: int = 2) -> int:
    """Mostra abas, cabeçalho e algumas linhas de um arquivo baixado."""
    try:
        import openpyxl
    except ImportError:
        print("Falta o openpyxl: pip install openpyxl")
        return 1
    import warnings

    warnings.filterwarnings("ignore")

    caminho = os.path.join(ITBI_DIR, f"itbi_{ano}.xlsx")
    if not os.path.exists(caminho):
        print(f"Não achei {caminho}. Rode --baixar --anos {ano} antes.")
        return 1

    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    print(f"{os.path.basename(caminho)}: {len(wb.sheetnames)} abas")
    print(f"  {', '.join(wb.sheetnames)}")

    nome_aba = aba or wb.sheetnames[0]
    if nome_aba not in wb.sheetnames:
        print(f"[aviso] aba {nome_aba!r} não existe; usando {wb.sheetnames[0]!r}")
        nome_aba = wb.sheetnames[0]

    ws = wb[nome_aba]
    print(f"\naba {nome_aba!r}: {ws.max_row} linhas x {ws.max_column} colunas")
    for i, linha in enumerate(ws.iter_rows(max_row=linhas + 1, values_only=True), 1):
        print(f"\n  -- linha {i} --")
        for j, v in enumerate(linha):
            if v is not None:
                print(f"     col {j:2d}: {str(v)[:58]}")
    wb.close()
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Base de ITBI da Prefeitura de SP")
    parser.add_argument("--baixar", action="store_true", help="baixa os arquivos")
    parser.add_argument("--listar", action="store_true", help="só lista os links")
    parser.add_argument("--inspecionar", action="store_true",
                        help="mostra a estrutura de um arquivo já baixado")
    parser.add_argument("--anos", nargs="+", type=int, metavar="ANO",
                        help="limita a estes anos (padrão: todos)")
    parser.add_argument("--ano", type=int, default=2026,
                        help="ano a inspecionar (padrão 2026)")
    parser.add_argument("--aba", default=None, help="aba a inspecionar")
    parser.add_argument("--forcar", action="store_true",
                        help="rebaixa mesmo se o arquivo já existir")
    args = parser.parse_args()

    if args.inspecionar:
        return inspecionar(args.ano, args.aba)
    if args.listar or args.baixar:
        if args.listar:
            links = descobrir_links()
            for ano, u in sorted(links.items()):
                print(f"  {ano}  {u.rsplit('/', 1)[-1][:66]}")
            if not args.baixar:
                return 0
        return baixar(args.anos, args.forcar)

    parser.print_help()
    return 0


if __name__ == "__main__":
    sys.exit(main())
